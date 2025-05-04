from kafka import KafkaConsumer
import PyPDF2
import docx
import os
import json
import mysql.connector
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
import logging
from tqdm import tqdm
import sys
import socketio
import requests
import time
import ollama
from dotenv import load_dotenv

load_dotenv()

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    filename='resume_processor.log',
    format='%(asctime)s %(levelname)s:%(message)s'
)

# Kafka Configuration
KAFKA_TOPIC = "resume-topic"
KAFKA_SERVER = "localhost:9092"

# SocketIO client
sio = socketio.Client()

try:
    sio.connect('http://localhost:5001')
    logging.info("Connected to SocketIO server")
except Exception as e:
    logging.error(f"Failed to connect to SocketIO server: {e}")

try:
    consumer = KafkaConsumer(
        KAFKA_TOPIC,
        bootstrap_servers=KAFKA_SERVER,
        auto_offset_reset="earliest",
        group_id="resume-consumer-group",
        value_deserializer=lambda x: json.loads(x.decode("utf-8")),
        enable_auto_commit=False
    )
    logging.info(f"Connected to Kafka broker, subscribed to topic: {KAFKA_TOPIC}")
except Exception as e:
    logging.error(f"Failed to initialize Kafka consumer: {e}")
    raise

def extract_text(file_path):
    """Extract text from a resume (PDF/DOCX/TXT)."""
    try:
        if file_path.endswith(".pdf"):
            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
                return text.strip()

        elif file_path.endswith(".docx"):
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            return text.strip()

        elif file_path.endswith(".txt"):
            with open(file_path, "r", encoding="utf-8") as file:
                return file.read().strip()

        else:
            logging.error(f"Unsupported file format: {file_path}")
            return None
    except Exception as e:
        logging.error(f"Error extracting text from {file_path}: {e}")
        return None

def extract_json(response_text):
    """Extract and validate JSON from response text."""
    logging.info(f"Raw API response: {response_text[:1000]}")
    brace_count = 0
    start = -1
    for i, char in enumerate(response_text):
        if char == '{':
            if start == -1:
                start = i
            brace_count += 1
        elif char == '}':
            brace_count -= 1
            if brace_count == 0 and start != -1:
                json_text = response_text[start:i + 1]
                try:
                    parsed = json.loads(json_text)
                    required_fields = ["name", "phone", "email", "skills", "job_predictions"]
                    if not all(field in parsed for field in required_fields):
                        return {
                            "error": f"Missing required fields: {', '.join(f for f in required_fields if f not in parsed)}",
                            "raw_output": response_text,
                            "extracted_text": json_text
                        }
                    # Validate job_predictions structure
                    for job in parsed.get("job_predictions", []):
                        job_required_fields = ["job", "score", "skills", "Reason"]
                        if not all(field in job for field in job_required_fields):
                            return {
                                "error": f"Missing required fields in job_predictions: {', '.join(f for f in job_required_fields if f not in job)}",
                                "raw_output": response_text,
                                "extracted_text": json_text
                            }
                    return parsed
                except json.JSONDecodeError as e:
                    return {
                        "error": f"Failed to parse extracted JSON: {e}",
                        "raw_output": response_text,
                        "extracted_text": json_text
                    }
    if start != -1 and brace_count > 0:
        json_text = response_text[start:] + '}'
        logging.info(f"Attempted fix by adding closing brace: {json_text[:1000]}")
        try:
            parsed = json.loads(json_text)
            required_fields = ["name", "phone", "email", "skills", "job_predictions"]
            if not all(field in parsed for field in required_fields):
                return {
                    "error": f"Missing required fields: {', '.join(f for f in required_fields if f not in parsed)}",
                    "raw_output": response_text,
                    "attempted_text": json_text
                }
            for job in parsed.get("job_predictions", []):
                job_required_fields = ["job", "score", "skills", "Reason"]
                if not all(field in job for field in job_required_fields):
                    return {
                        "error": f"Missing required fields in job_predictions: {', '.join(f for f in job_required_fields if f not in job)}",
                        "raw_output": response_text,
                        "attempted_text": json_text
                    }
            return parsed
        except json.JSONDecodeError as e:
            return {
                "error": f"Failed to parse after fix attempt: {e}",
                "raw_output": response_text,
                "attempted_text": json_text
            }
    return {
        "error": "No complete JSON object found",
        "raw_output": response_text
    }

def process_resume_with_llama(resume_text):
    """Predict up to 3 job roles with confidence scores, skills, and reasons using LLaMA."""
    prompt = f"""
Extract details from the resume and predict up to 3 job roles with confidence scores, relevant skills, and reasons for the scores.

### *STRICT RULES:*
- *Output ONLY JSON.* No explanations or extra text.
- *Ensure valid JSON.* No missing brackets or incorrect structures.
- *Follow the exact format* without placeholders.
- If any data is missing, use "None" for strings or [] for lists.
- *job_predictions* must be a list of up to 3 objects with job, score, skills, and Reason.

### *Data Extraction Requirements:*
- *name:* Extract the full name accurately.
- *phone:* Extract a valid phone number as a string (or "None").
- *email:* Extract a valid email address as a string (or "None").
- *skills:* Extract all skills listed in the resume as a comma-separated string.
- *job_predictions:* Predict up to 3 job roles from the following list: [Data Scientist, Software Engineer, Product Manager, DevOps Engineer, UX Designer]. Each prediction includes:
  - *job:* The job role.
  - *score:* Confidence score (0-100) based on skills, experience, and education.
  - *skills:* List of resume skills relevant to this job.
  - *Reason:* A brief explanation of the score based on skills match (50%), experience (30%), and education/certifications (20%).

### *Prediction Guidelines:*
- *Skills Match (50%)* → Match resume skills to job requirements.
- *Experience (30%)* → Evaluate relevant work/projects.
- *Education/Certifications (20%)* → Consider relevant qualifications.
- *Confidence Scores:* Sum to 100% across predicted jobs if multiple, or 100% for one job.
- *Skills per Job:* Only include skills from the resume that are relevant to the job.
- *Reason for Score:* Explain the score breakdown like why you given respective marks. More like a positives of the applicant".

### *Job Role Skill Requirements:*
- *Data Scientist:* Python, R, SQL, machine learning, statistics, data analysis.
- *Software Engineer:* Python, Java, JavaScript, C++, git, agile.
- *Product Manager:* Product management, agile, UX, stakeholder management.
- *DevOps Engineer:* Docker, Kubernetes, CI/CD, AWS, Linux.
- *UX Designer:* UX design, Figma, user research, prototyping.

Resume Content:
{resume_text}

### *Expected JSON Output:*
{{
  "name": "<Full Name or 'None'>",
  "phone": "<Phone Number or 'None'>",
  "email": "<Email or 'None'>",
  "skills": "<Comma-separated skills or 'None'>",
  "job_predictions": [
    {{"job": "<Job Role>", "score": <Score 0-100>, "skills": ["<Skill1>", "<Skill2>"], "Reason": "<Explanation of score>"}},
    {{"job": "<Job Role>", "score": <Score 0-100>, "skills": ["<Skill1>", "<Skill2>"], "Reason": "<Explanation of score>"}},
    {{"job": "<Job Role>", "score": <Score 0-100>, "skills": ["<Skill1>", "<Skill2>"], "Reason": "<Explanation of score>"}}
  ]
}}
    """

    url = "https://shaggy-frogs-fold.loca.lt/api/generate"
    payload = {
        "model": "llama3.1:8b",
        "prompt": prompt
    }
    max_retries = 3
    retry_delay = 5

    for attempt in range(max_retries):
        try:
            logging.info(f"Sending request to API (attempt {attempt + 1}/{max_retries})")
            response = requests.post(
                url,
                json=payload,
                headers={
                    "Content-Type": "application/json",
                    "bypass-tunnel-reminder": "true"
                },
                timeout=60,
                stream=True
            )
            response.raise_for_status()

            raw_output = ""
            for line in response.iter_lines(decode_unicode=True):
                if line:
                    try:
                        json_line = json.loads(line)
                        if "response" in json_line:
                            raw_output += json_line["response"]
                        if json_line.get("done", False):
                            break
                    except json.JSONDecodeError as e:
                        logging.error(f"Failed to parse NDJSON line: {line}, Error: {e}")
                        continue

            if not raw_output:
                logging.error("Empty response from API")
                return {"error": "Empty response from API", "raw_output": ""}

            logging.info("Received response from API")
            extracted_data = extract_json(raw_output)
            return extracted_data
        except requests.exceptions.RequestException as e:
            logging.error(f"API error (attempt {attempt + 1}): {str(e)}")
            if attempt < max_retries - 1:
                logging.info(f"Retrying after {retry_delay} seconds...")
                time.sleep(retry_delay)
            else:
                logging.error("Max retries reached. Falling back to local Ollama.")
                try:
                    logging.info("Sending request to local Ollama server")
                    ollama_response = ollama.chat(
                        model="llama3.2",
                        messages=[{"role": "user", "content": prompt}]
                    )
                    raw_output = ollama_response["message"]["content"]
                    logging.info("Received response from Ollama")
                    return extract_json(raw_output)
                except Exception as e2:
                    logging.error(f"Local Ollama error: {str(e2)}")
                    return {"error": f"API error: {str(e)}, Local error: {str(e2)}", "raw_output": ""}
        except ValueError as e:
            logging.error(f"Response parsing error: {str(e)}")
            return {"error": f"Invalid response: {str(e)}", "raw_output": ""}

def save_to_database(resume_data, file_path):
    try:
        conn = mysql.connector.connect(
            host=os.getenv("host"),
            user=os.getenv("user"),
            password=os.getenv("password"),
            database=os.getenv("database"),
        )
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM resumes WHERE file_path = %s", (file_path,))
        if cursor.fetchone():
            logging.info(f"Skipped duplicate resume: {file_path}")
            sio.emit('resume_processed', {
                'file_path': file_path,
                'status': 'failed',
                'error': 'Duplicate resume',
                'filename': os.path.basename(file_path)
            })
            return

        # Insert into resumes table
        query = """
        INSERT INTO resumes (name, phone, email, skills, uploaded_at, file_path)
        VALUES (%s, %s, %s, %s, %s, %s)
        """
        values = (
            resume_data.get("name", "None"),
            resume_data.get("phone", "None"),
            resume_data.get("email", "None"),
            resume_data.get("skills", "None"),
            datetime.now(),
            file_path
        )
        cursor.execute(query, values)
        resume_id = cursor.lastrowid

        # Insert job predictions into resume_jobs table
        job_predictions = resume_data.get("job_predictions", [])
        if job_predictions:
            job_query = """
            INSERT INTO resume_jobs (resume_id, job_role, score, job_skills, Reason)
            VALUES (%s, %s, %s, %s, %s)
            """
            for job in job_predictions:
                job_values = (
                    resume_id,
                    job.get("job", "Unknown"),
                    job.get("score", 0),
                    ",".join(job.get("skills", [])) if job.get("skills") else "None",
                    job.get("Reason", "No reason provided")
                )
                cursor.execute(job_query, job_values)

        conn.commit()
        logging.info(f"Successfully saved resume: {file_path}")
        cursor.close()
        conn.close()

        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'completed',
            'data': resume_data,
            'filename': os.path.basename(file_path)
        })
    except mysql.connector.Error as e:
        logging.error(f"Database error while saving resume: {file_path}, Error: {e}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': str(e),
            'filename': os.path.basename(file_path)
        })

def process_message(message, consumer):
    logging.info(f"Received Kafka message: {message.value}")
    message_data = message.value
    file_path = message_data.get("file_path", "")
    filename = message_data.get("original_filename", os.path.basename(file_path))

    if not os.path.exists(file_path):
        logging.error(f"File not found: {file_path}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': 'File not found',
            'filename': filename
        })
        try:
            consumer.commit()
            logging.info(f"Committed Kafka offset for message: {message.offset}")
        except Exception as e:
            logging.error(f"Failed to commit Kafka offset: {e}")
        return

    logging.info(f"Started processing: {file_path}")
    sio.emit('resume_processing', {
        'file_path': file_path,
        'status': 'processing',
        'filename': filename
    })

    resume_text = extract_text(file_path)

    if resume_text:
        extracted_details = process_resume_with_llama(resume_text)
        logging.info(f"Extracted details for: {file_path}")
        save_to_database(extracted_details, file_path)
    else:
        logging.error(f"Failed to extract text from: {file_path}")
        sio.emit('resume_processed', {
            'file_path': file_path,
            'status': 'failed',
            'error': 'Failed to extract text',
            'filename': filename
        })

    try:
        consumer.commit()
        logging.info(f"Committed Kafka offset for message: {message.offset}")
    except Exception as e:
        logging.error(f"Failed to commit Kafka offset: {e}")

logging.info("Starting Kafka consumer loop")
with ThreadPoolExecutor(max_workers=3) as executor:
    progress_bar = tqdm(desc="Processed Resumes", unit="resume", file=sys.stdout)
    futures = []
    batch_size = 3
    for message in consumer:
        logging.info("Fetched a message from Kafka")
        futures.append(executor.submit(process_message, message, consumer))
        progress_bar.update(1)
        if len(futures) >= batch_size:
            for future in futures:
                future.result()
            futures = []
    for future in futures:
        future.result()
logging.info("Consumer loop ended")